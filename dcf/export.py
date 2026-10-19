"""Excel output: one sheet per step of the model, plus charts."""
from __future__ import annotations

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

from .valuation import Assumptions

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MILLIONS = '#,##0.0,,"M"'
PCT = "0.0%"


def _style(ws, number_format: str | None = None, first_col_width: int = 28):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.column_dimensions["A"].width = first_col_width
    for col in ws.iter_cols(min_col=2):
        ws.column_dimensions[col[0].column_letter].width = 16
        if number_format:
            for c in col[1:]:
                c.number_format = number_format
    ws.freeze_panes = "B2"


def to_excel(path: str, ticker: str, hist: pd.DataFrame, a: Assumptions,
             cost_of_capital: dict, result: dict, sens: pd.DataFrame,
             hist_fcff: pd.Series, sources: dict | None = None) -> None:
    summary = pd.DataFrame([
        ("Company", ticker),
        ("WACC", result["wacc"]),
        ("Terminal growth", result["terminal_growth"]),
        ("PV of forecast FCFF", result["pv_forecast_fcff"]),
        ("PV of terminal value", result["pv_terminal_value"]),
        ("Enterprise value", result["enterprise_value"]),
        ("(-) Total debt", result["total_debt"]),
        ("(+) Cash", result["cash"]),
        ("Equity value", result["equity_value"]),
        ("Shares", result["shares"]),
        ("Value per share", result["value_per_share"]),
        ("Terminal value as % of EV", result["terminal_share_of_ev"]),
    ], columns=["item", "value"])

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        hist.T.to_excel(xw, sheet_name="Historicals")
        result["projection"].T.to_excel(xw, sheet_name="Projection")
        a.to_frame().to_excel(xw, sheet_name="Assumptions", index=False)
        pd.DataFrame(cost_of_capital.items(), columns=["component", "value"]) \
            .to_excel(xw, sheet_name="WACC", index=False)
        sens.to_excel(xw, sheet_name="Sensitivity")
        if sources:
            pd.DataFrame(sources.items(), columns=["assumption", "source and reasoning"]) \
                .to_excel(xw, sheet_name="Sources", index=False)

        wb = xw.book
        ws = wb["Summary"]
        _style(ws)
        pct_rows = {"WACC", "Terminal growth", "Terminal value as % of EV"}
        for row in ws.iter_rows(min_row=2):
            label, cell = row[0].value, row[1]
            if label in pct_rows:
                cell.number_format = PCT
            elif label == "Value per share":
                cell.number_format = '"$"#,##0.00'
            elif label == "Shares":
                cell.number_format = '#,##0.0,,"M"'
            elif isinstance(cell.value, (int, float)):
                cell.number_format = MILLIONS
        _style(wb["Historicals"], MILLIONS)
        _style(wb["Projection"], MILLIONS)
        for r in wb["Projection"].iter_rows(min_row=2):
            if r[0].value == "discount_factor":
                for c in r[1:]:
                    c.number_format = "0.0000"
        _style(wb["Assumptions"])
        _style(wb["WACC"], PCT)
        _style(wb["Sensitivity"], '"$"#,##0.00', first_col_width=14)
        ws_s = wb["Sensitivity"]
        ws_s.conditional_formatting.add(
            f"B2:{ws_s.cell(ws_s.max_row, ws_s.max_column).coordinate}",
            ColorScaleRule(start_type="min", start_color="CDE2FB",
                           end_type="max", end_color="256ABF"))
        _charts_sheet(wb, hist, result, hist_fcff)
        if sources:
            ws_src = wb["Sources"]
            _style(ws_src, first_col_width=26)
            ws_src.column_dimensions["B"].width = 120
            for c in ws_src["B"][1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")


def _charts_sheet(wb, hist: pd.DataFrame, result: dict, hist_fcff: pd.Series) -> None:
    # small table on the sheet + native Excel bar charts that read from it
    ws = wb.create_sheet("Charts", 1)
    ws.append(["Year", "Revenue (hist)", "Revenue (proj)", "FCFF (hist)", "FCFF (proj)"])
    proj = result["projection"]
    for y in hist.index:
        ws.append([str(y), hist.loc[y, "revenue"] / 1e9, None, hist_fcff[y] / 1e9, None])
    for y in proj.index:
        ws.append([str(y), None, proj.loc[y, "revenue"] / 1e9, None, proj.loc[y, "fcff"] / 1e9])
    _style(ws, '#,##0.00', first_col_width=10)
    last = ws.max_row
    cats = Reference(ws, min_col=1, min_row=2, max_row=last)

    for title, first_col, anchor in (("Revenue ($B)", 2, "G2"),
                                     ("Unlevered free cash flow ($B)", 4, "G20")):
        chart = BarChart()
        chart.title, chart.type, chart.overlap, chart.gapWidth = title, "col", 100, 40
        chart.height, chart.width = 8, 16
        chart.add_data(Reference(ws, min_col=first_col, max_col=first_col + 1,
                                 min_row=1, max_row=last), titles_from_data=True)
        chart.set_categories(cats)
        for series, color in zip(chart.series, ("2A78D6", "EB6834")):
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.noFill = True
        chart.legend.position = "b"
        ws.add_chart(chart, anchor)
